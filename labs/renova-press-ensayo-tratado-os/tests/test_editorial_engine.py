from pathlib import Path
import importlib.util
import json
import subprocess
import sys
import pytest
pytest.importorskip('docx')
from docx import Document

HERE=Path(__file__).resolve().parents[1]
CLI=HERE/'scripts'/'editorial_engine.py'

def fixture(tmp_path):
    p=tmp_path/'test_ensayo.docx'
    d=Document();d.add_heading('Tesis y contratesis',1)
    d.add_paragraph(('La vida exige una crítica fundada, pero no toda crítica es verdadera. ' * 17).strip())
    d.add_paragraph(('La segunda premisa conserva el retorno a la cuestión original. ' * 12).strip())
    d.add_paragraph('Bibliografía\nObra de prueba. Sin referencias externas.');d.save(p)
    return p

def call(*args):
    p=subprocess.run([sys.executable,str(CLI),*map(str,args)],capture_output=True,text=True)
    assert p.returncode==0,(p.stdout,p.stderr)
    return p.stdout

def test_index_preserves_source(tmp_path):
    f=fixture(tmp_path); content=f.read_bytes();out=tmp_path/'reg.json'
    call('index',f,'--out',out)
    r=json.loads(out.read_text()); assert r['items'][0]['status']=='INGESTED_NOT_VERIFIED'
    assert r['items'][0]['source'] is None
    assert f.read_bytes()==content

def test_style_and_audit(tmp_path):
    f=fixture(tmp_path);out=tmp_path/'style.json';audit=tmp_path/'audit.json'
    call('style',f,'--out',out); call('audit',f,'--out',audit)
    assert json.loads(out.read_text())['corpus_files']==1
    assert json.loads(audit.read_text())['status']=='REQUIRES_REVISION_HUMANA'

def test_typeset_keeps_text_and_justifies(tmp_path):
    f=fixture(tmp_path); out=tmp_path/'proof.docx'
    call('typeset',f,'--out',out)
    orig=Document(f); new=Document(out)
    assert [p.text for p in orig.paragraphs]==[p.text for p in new.paragraphs]
    assert abs(new.sections[0].page_width.cm-15.24)<0.02
    assert abs(new.sections[0].page_height.cm-22.86)<0.02
    assert f.exists() and out.exists()

def test_cannot_overwrite_original(tmp_path):
    f=fixture(tmp_path)
    p=subprocess.run([sys.executable,str(CLI),'typeset',str(f),'--out',str(f)],capture_output=True,text=True)
    assert p.returncode!=0
