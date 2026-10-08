#!/usr/bin/env python3
"""ℛenova Press · Essays/Treatises OS. Local, offline, review-first CLI.

No external API calls, credentials, or silent publishing. This is a production
scaffold and a conservative manuscript handling pipeline, not a fact checker.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from statistics import median

EXTENSIONS={'.docx','.md','.txt','.pdf'}

def sha256(path:Path)->str:
    h=hashlib.sha256()
    with path.open('rb') as f:
        for part in iter(lambda: f.read(1<<20), b''): h.update(part)
    return h.hexdigest()

def extract_paragraphs(path:Path):
    """Return semantic paragraphs; never silently overwrite the source."""
    if path.suffix.lower()=='.docx':
        from docx import Document
        d=Document(str(path))
        return [{'text':p.text, 'style':p.style.name if p.style else 'Normal'}
                for p in d.paragraphs if p.text.strip()]
    if path.suffix.lower()=='.pdf':
        import fitz
        with fitz.open(str(path)) as pdf:
            # PDF extraction has no trustworthy semantic paragraph structure.
            blocks=[]
            for pg in pdf:
                for b in pg.get_text('blocks'):
                    if len(b)>4 and b[4].strip():
                        blocks.append({'text':b[4].strip(),'style':'PDF_BLOCK_UNVERIFIED'})
            return blocks
    text=path.read_text(encoding='utf-8-sig')
    groups=re.split(r'\n\s*\n',text)
    return [{'text':g.strip(),'style':'Heading' if g.strip().startswith('#') else 'Normal'}
            for g in groups if g.strip()]

def metrics(path:Path):
    ps=extract_paragraphs(path)
    tokens=[len(re.findall(r'\b[\wÀ-ÿ]+\b',p['text'])) for p in ps]
    text='\n'.join(p['text'] for p in ps)
    return {
       'sha256':sha256(path),'format':path.suffix.lower(),
       'paragraphs':len(ps),'words':sum(tokens),'median_words_per_paragraph':median(tokens) if tokens else 0,
       'long_paragraphs_100w':sum(x>=100 for x in tokens),
       'parenthetical_asides':text.count('('),
       'interrogations':text.count('¿'),'em_dashes':text.count('—'),
       'quotation_markers':sum(text.count(c) for c in ['«','»','“','”','"']),
       'headings_detected':sum(p['style'].lower().startswith(('heading','título','titulo')) for p in ps),
       'extraction_warning':'PDF blocks do not equal author paragraphs' if path.suffix.lower()=='.pdf' else None,
    }

def cmd_index(a):
    source=Path(a.source).expanduser().resolve()
    if not source.exists(): raise FileNotFoundError(source)
    paths=([source] if source.is_file() else sorted(p for p in source.rglob('*') if p.is_file() and p.suffix.lower() in EXTENSIONS))
    records=[]
    for p in paths:
        if p.suffix.lower() not in EXTENSIONS:continue
        try:
            m=metrics(p)
            records.append({'filename':p.name,'source':str(p) if a.store_paths else None,**m,'status':'INGESTED_NOT_VERIFIED'})
        except Exception as ex: records.append({'filename':p.name,'status':'ERROR','error':type(ex).__name__+': '+str(ex)})
    doc={'schema_version':'1.0','created_utc':datetime.now(timezone.utc).isoformat(),
         'privacy':'No manuscript contents retained in this index. Paths omitted unless --store-paths.',
         'items':records}
    dst=Path(a.out);dst.parent.mkdir(parents=True,exist_ok=True);dst.write_text(json.dumps(doc,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'items':len(records),'success':sum(i['status']!='ERROR' for i in records),'output':str(dst)},ensure_ascii=False))

def cmd_style(a):
    src=Path(a.source)
    records=[]
    for p in (sorted(src.rglob('*.docx')) if src.is_dir() else [src]):
        try:records.append({'file':p.name,**metrics(p)})
        except Exception as ex: records.append({'file':p.name,'error':str(ex)})
    valid=[r for r in records if 'words' in r]
    profile={
      'status':'DESCRIPTIVE_NOT_PRESCRIPTIVE','scope':'Archival style metrics; not a license to fabricate quotes',
      'corpus_files':len(valid),
      'total_words':sum(r['words'] for r in valid),
      'median_words_per_paragraph_across_files':median([r['median_words_per_paragraph'] for r in valid]) if valid else 0,
      'characteristics':['Ensayo dialógico y polémico','Apertura problemática o narrativa','Digresiones históricas con retorno a tesis','Alternancia de sátira, erudición y primera persona','Párrafos extensos cuando el argumento los requiere','Citas a verificar en obra y edición original'],
      'source_metrics':records}
    dst=Path(a.out);dst.parent.mkdir(parents=True,exist_ok=True);dst.write_text(json.dumps(profile,indent=2,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'files':len(valid),'output':str(dst)},ensure_ascii=False))

def _add_page_field(run):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)

def cmd_typeset(a):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    path=Path(a.source)
    if path.suffix.lower()=='.docx': d=Document(str(path))
    elif path.suffix.lower() in {'.md','.txt'}:
        d=Document()
        for item in extract_paragraphs(path):
            s=item['text']
            if item['style']=='Heading':d.add_heading(s.lstrip('#').strip(),level=1)
            else:d.add_paragraph(s.replace('\n',' '))
    else:raise ValueError('typeset accepts DOCX, TXT or MD; PDF extraction requires manual structural validation')
    for s in d.sections:
        s.page_width=Cm(15.24);s.page_height=Cm(22.86) # 6 x 9 in
        s.top_margin=Cm(2.20);s.bottom_margin=Cm(2.25)
        s.left_margin=Cm(2.40);s.right_margin=Cm(2.00)
        s.header_distance=Cm(1.25);s.footer_distance=Cm(1.25)
        if not s.different_first_page_header_footer:
            s.different_first_page_header_footer=True
        # Remove stale manuscript headers like “Página 3 de 5” after re-pagination.
        # The freshly composed proof instead carries a live folio in the footer.
        for hp in s.header.paragraphs:
            h=hp.text.lower()
            if ('página' in h or 'page' in h) and ' de ' in h or ('page' in h and ' of ' in h):
                hp.text=''
        f=s.footer
        if not any('PAGE' in x._p.xml for x in f.paragraphs):
            p=f.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            _add_page_field(p.add_run())
    styles=d.styles
    body=styles['Normal'];body.font.name='Liberation Serif';body.font.size=Pt(11.2)
    body.paragraph_format.line_spacing=1.14
    body.paragraph_format.space_after=Pt(0)
    body.paragraph_format.space_before=Pt(0)
    body.paragraph_format.first_line_indent=Cm(0.48)
    # Preserve all existing runs, emphasis, heading styles, tables and images.
    for p in d.paragraphs:
        name=p.style.name.lower() if p.style else ''
        if name.startswith(('heading','título','titulo')):
            p.paragraph_format.space_before=Pt(12)
            p.paragraph_format.space_after=Pt(7)
            p.paragraph_format.keep_with_next=True
        elif p.text.strip():
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.widow_control=True
    dest=Path(a.out);dest.parent.mkdir(parents=True,exist_ok=True)
    if dest.resolve()==path.resolve():raise ValueError('Never overwrite original manuscript')
    d.save(str(dest))
    print(json.dumps({'proof_docx':str(dest),'warning':'Proof only. Verify citations, footnotes, and every rendered page before publication.'},ensure_ascii=False))

def cmd_render(a):
    inp=Path(a.source);out=Path(a.out_dir)
    if inp.suffix.lower()!='.docx':raise ValueError('render expects DOCX')
    out.mkdir(parents=True,exist_ok=True)
    lo=shutil.which('libreoffice') or shutil.which('soffice')
    if not lo:raise EnvironmentError('LibreOffice is needed for local PDF render')
    # Isolate conversion profile; no network use.
    profile=(out/'.lo-profile').resolve().as_uri()
    result=subprocess.run([lo,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(out),str(inp)],capture_output=True,text=True,timeout=120)
    pdf=out/(inp.stem+'.pdf')
    if result.returncode!=0 or not pdf.exists():raise RuntimeError(result.stdout+'\n'+result.stderr)
    import fitz
    with fitz.open(str(pdf)) as p:
        pages=len(p)
        empty=[i+1 for i,pg in enumerate(p) if not pg.get_text().strip()]
        overflow=[]
        for i,pg in enumerate(p):
            rect=pg.rect
            for b in pg.get_text('blocks'):
                if len(b)>4 and b[4].strip() and (b[0]<-0.5 or b[1]<-0.5 or b[2]>rect.width+0.5 or b[3]>rect.height+0.5):
                    overflow.append(i+1);break
    report={'pdf':str(pdf),'pages':pages,'fully_empty_pages':empty,'out_of_page_text_boxes':overflow,
            'result':'STRUCTURAL_CHECK_ONLY','manual_gate':'Visual inspect every page; typography and source fidelity not checked.'}
    (out/'preflight.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))

def cmd_audit(a):
    p=Path(a.source); ps=extract_paragraphs(p)
    words=[len(re.findall(r'\b[\wÀ-ÿ]+\b',x['text'])) for x in ps]
    issues=[]; fatal=[]
    if p.suffix.lower()=='.pdf':issues.append('PDF text blocks cannot certify author paragraph fidelity.')
    if sum(words)<250:issues.append('Short text: confirm this is an essay, rather than a full monograph.')
    if p.suffix.lower()=='.docx':
        import zipfile
        with zipfile.ZipFile(p) as z:
            if 'word/footnotes.xml' in z.namelist():issues.append('Footnotes present: manually verify placement, numbering and complete survival after export.')
    text='\n'.join(x['text'] for x in ps)
    for label,pattern in [('reference placeholder',r'\[(?:DOI|CITA|FUENTE|REF|PÁGINA|PAGE)[^\]]*\]'),('draft placeholder',r'\b(?:TODO|TBD|LOREM IPSUM)\b')]:
        hits=len(re.findall(pattern,text,re.I))
        if hits:issues.append(f'{hits} possible {label}(s)')
    if not re.search(r'\b(bibliograf[ií]a|referencias|works cited|references)\b',text,re.I):
        issues.append('No bibliography heading detected; check whether this genre requires references.')
    gate={
      'source':p.name,'sha256':sha256(p),'total_words':sum(words),'paragraphs':len(ps),
      'automated_findings':issues,'blocking_issues':fatal,'status':'REQUIRES_REVISION_HUMANA',
      'manual_gates':['Author approves thesis and voice','Source-by-source citation and quotation check',
                      'Editorial proofreading complete','Layout proof per page',
                      'Permissions/right-to-reproduce verified','ISBN and legal notices verified if needed',
                      'Digital accessibility and commercial formats verified']}
    dest=Path(a.out);dest.parent.mkdir(parents=True,exist_ok=True);dest.write_text(json.dumps(gate,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(dest),'status':gate['status'],'findings':len(issues)},ensure_ascii=False))

def main():
    parser=argparse.ArgumentParser(description='ℛenova Press · Ensayo/Tratado OS · offline, review-first')
    sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('index');p.add_argument('source');p.add_argument('--out',required=True);p.add_argument('--store-paths',action='store_true');p.set_defaults(func=cmd_index)
    p=sub.add_parser('style');p.add_argument('source');p.add_argument('--out',required=True);p.set_defaults(func=cmd_style)
    p=sub.add_parser('typeset');p.add_argument('source');p.add_argument('--out',required=True);p.set_defaults(func=cmd_typeset)
    p=sub.add_parser('render');p.add_argument('source');p.add_argument('--out-dir',required=True);p.set_defaults(func=cmd_render)
    p=sub.add_parser('audit');p.add_argument('source');p.add_argument('--out',required=True);p.set_defaults(func=cmd_audit)
    a=parser.parse_args();a.func(a)
if __name__=='__main__':main()
